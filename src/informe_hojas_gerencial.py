"""
informe_hojas_gerencial.py
--------------------------
Hojas orientadas al POM: Resumen ejecutivo (KPIs, hallazgos, decisiones) y Metodologia.
"""

from typing import Any

from informe_estilos import (FMT_ENTERO, FMT_PCT, GRIS_CLARO, ajustar_anchos, seccion,
                             tabla, titulo)
from openpyxl.styles import Alignment, Font, PatternFill


def hoja_resumen(wb, res: dict[str, Any]) -> None:
    """Hoja 1: lo que el POM necesita para decidir, en una sola pantalla."""
    ws = wb.active
    ws.title = "Resumen ejecutivo"
    titulo(ws, "Informe de cruce CSV vs ODS - Resumen ejecutivo",
           f"Generado: {res['generado']}  |  Base de comparacion: CSV de casos QA  |  IDs enmascarados")
    ws.sheet_view.showGridLines = False

    fila = seccion(ws, 4, "1. Indicadores clave (KPI)", 6)
    kpis = res["kpis"]
    filas = []
    for k in kpis:
        valor = k["valor"]
        filas.append([k["nombre"], valor, k["unidad"], k["semaforo"], k["umbral"], k["detalle"]])
    fila = tabla(ws, fila, ["Indicador", "Valor", "Unidad", "Estado", "Umbral", "Detalle"], filas,
                 formatos={2: "#,##0.##"}, col_estado=4)

    fila = seccion(ws, fila + 1, "2. Hallazgos principales", 6)
    for h in res["hallazgos"]:
        ws.cell(row=fila, column=1, value="- " + h).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        ws.row_dimensions[fila].height = 30 if len(h) > 110 else 18
        fila += 1

    fila = seccion(ws, fila + 1, "3. Decisiones que se solicitan", 6)
    filas = [[d["prioridad"], d["decision"], d["soporte"], d["responsable"]] for d in res["decisiones"]]
    fila = tabla(ws, fila, ["Prioridad", "Decision", "Soporte", "Responsable"], filas)

    fila = seccion(ws, fila + 1, "4. Alcance y advertencias", 6)
    s = res["resumen"]
    notas = [
        f"Se cruzaron {s['csv_total']} casos del CSV contra las {len(res['hojas_calidad'])} hojas del ODS "
        f"({s['ods_validos']} IDs validos).",
        "Cobertura QA y trazabilidad miden cosas distintas; ver hoja Metodologia.",
        "Los umbrales del semaforo son una propuesta tecnica pendiente de validacion por el POM.",
        "El analisis usa solo tipo e ID. Las demas columnas del ODS son auxiliares y no se interpretan.",
    ]
    for n in notas:
        ws.cell(row=fila, column=1, value="- " + n).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        ws.row_dimensions[fila].height = 18
        fila += 1

    ajustar_anchos(ws, minimo=12, maximo=70)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["F"].width = 52
    ws.freeze_panes = "A4"


def hoja_metodologia(wb, res: dict[str, Any]) -> None:
    """Definiciones, supuestos y limitaciones, para que el POM pueda defender los numeros."""
    ws = wb.create_sheet("Metodologia")
    titulo(ws, "Metodologia, definiciones y limitaciones")
    ws.sheet_view.showGridLines = False
    u = res["umbrales"]

    fila = seccion(ws, 3, "Definiciones de los indicadores", 3)
    filas = [
        ["Trazabilidad CSV a ODS",
         "Casos del CSV cuyo tipo e ID existen en la hoja ODS de ese tipo, dividido entre los casos del CSV.",
         "Responde: los casos de prueba, existen en el estimador?"],
        ["Cobertura QA del ODS",
         "IDs validos del ODS que tienen un caso en el CSV, dividido entre los IDs validos del ODS.",
         "Responde: que parte del estimador se esta probando?"],
        ["Calidad del ODS",
         "Filas validas entre filas candidatas. Candidata: fila con tipo o ID. No cuenta encabezados ni filas vacias.",
         "Mide el estado del archivo ODS, no de las pruebas."],
        ["Fila valida", "Tipo numerico e ID numerico distinto de 0.", ""],
        ["Coincidencia exacta", "Mismo tipo e ID, y la fila esta en la hoja cuyo prefijo es ese tipo.", ""],
    ]
    fila = tabla(ws, fila, ["Termino", "Definicion", "Interpretacion"], filas)

    fila = seccion(ws, fila + 1, "Umbrales del semaforo (propuesta)", 3)
    filas = [[n.replace("_", " "), f"VERDE >= {v['verde']:g}%", f"AMARILLO >= {v['amarillo']:g}%, ROJO por debajo"]
             for n, v in u.items()]
    fila = tabla(ws, fila, ["Indicador", "Verde", "Amarillo / Rojo"], filas)
    ws.cell(row=fila, column=1, value="No existe un criterio de aceptacion documentado. "
            "Se ajustan en config/configuracion.yaml.").font = Font(italic=True, color="9C0006")
    fila += 1

    fila = seccion(ws, fila + 1, "Reglas de lectura del ODS", 3)
    filas = [
        ["Clave de cruce", "Columna 1 = tipo de documento, columna 2 = numero de ID. Las demas columnas son auxiliares.", ""],
        ["Normalizacion", "Se eliminan espacios, comillas y guiones del ID antes de comparar.", ""],
        ["Notacion cientifica",
         "Si el ODS muestra un ID como 1,23E+09, se usa el valor numerico almacenado en el archivo.",
         f"{res['resumen']['notacion_cientifica']} ID(s) recuperado(s)"],
        ["Filas excluidas",
         "ID igual a 0, tipo sin ID, ID sin tipo y formatos no reconocidos. Se listan en la hoja Excluidos ODS.", ""],
        ["Cruce contra todas las hojas",
         "Cada caso del CSV se busca en las 10 hojas, no solo en la de su tipo; asi se detectan IDs ubicados en otra hoja.", ""],
    ]
    fila = tabla(ws, fila, ["Regla", "Descripcion", "Observacion"], filas)

    fila = seccion(ws, fila + 1, "Limitaciones", 3)
    limites = [
        "El analisis valida existencia de tipo e ID. No evalua si el resultado del estimador de ingresos es correcto.",
        "Una cobertura alta no prueba que los casos sean suficientes ni representativos.",
        "Un ID sin caso QA no es un defecto por si mismo; es una decision de alcance.",
        "Las filas excluidas del ODS pueden contener registros reales que merecen revision en el origen.",
        "El informe refleja el estado de los archivos a la fecha de generacion.",
    ]
    for texto in limites:
        ws.cell(row=fila, column=1, value="- " + texto).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=3)
        ws.row_dimensions[fila].height = 18
        fila += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 95
    ws.column_dimensions["C"].width = 50
