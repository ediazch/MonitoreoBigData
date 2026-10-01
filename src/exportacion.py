"""
exportacion.py
--------------
Exporta el resultado del cruce ODS vs CSV a un archivo Excel (.xlsx).

Seguridad:
    - Los num_id se enmascaran siempre (106****185).
    - El archivo se guarda en data/processed/, nunca en la carpeta de datos originales.
"""

import os
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_ENCABEZADO = "1F4E78"
COLOR_OK = "C6EFCE"
COLOR_ALERTA = "FFEB9C"
COLOR_CRITICO = "FFC7CE"


def enmascarar_id(num_id: str) -> str:
    """Enmascara un ID dejando visibles los 3 primeros y 3 ultimos digitos."""
    if len(num_id) <= 6:
        return "***"
    return num_id[:3] + "*" * (len(num_id) - 6) + num_id[-3:]


def _escribir_hoja(ws, encabezados: list[str], filas: list[list[Any]]) -> None:
    """Escribe encabezados con estilo, filas de datos y ajusta anchos de columna."""
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in filas:
        ws.append(fila)

    for idx, titulo in enumerate(encabezados, start=1):
        ancho = max(
            [len(str(titulo))] + [len(str(f[idx - 1])) for f in filas if idx - 1 < len(f)]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(ancho + 3, 60)

    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions


def _color_por_cobertura(pct: float) -> str:
    if pct >= 99:
        return COLOR_OK
    if pct >= 90:
        return COLOR_ALERTA
    return COLOR_CRITICO


def exportar_excel(resultado: dict, stats: dict, directorio_salida: str) -> str:
    """
    Genera el Excel con 5 hojas: Resumen, Por tipo, Coincidencias, Solo CSV, Solo ODS.

    Args:
        resultado: Salida de transformacion.cruzar_datos().
        stats: Salida de transformacion.calcular_estadisticas().
        directorio_salida: Carpeta destino (se crea si no existe).

    Returns:
        Ruta absoluta del archivo generado.
    """
    os.makedirs(directorio_salida, exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    ruta = os.path.join(directorio_salida, f"analisis_relacion_{marca}.xlsx")

    wb = Workbook()

    # Hoja 1: resumen global
    ws = wb.active
    ws.title = "Resumen"
    _escribir_hoja(
        ws,
        ["Metrica", "Valor"],
        [
            ["Fecha de generacion", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Registros CSV", stats["total_registros_csv"]],
            ["Registros ODS", stats["total_registros_ods"]],
            ["Coincidencias", stats["coincidencias"]],
            ["Solo en CSV", stats["solo_en_csv"]],
            ["Solo en ODS", stats["solo_en_ods"]],
            ["Cobertura CSV (%)", stats["cobertura_csv_pct"]],
            ["Cobertura ODS (%)", stats["cobertura_ods_pct"]],
        ],
    )

    # Hoja 2: resumen por tipo con semaforo de cobertura
    ws = wb.create_sheet("Por tipo")
    filas_tipo = [
        [
            tipo,
            d["descripcion"],
            d["total_csv"],
            d["total_ods"],
            d["coincidencias"],
            d["cobertura_pct"],
        ]
        for tipo, d in sorted(resultado["resumen_por_tipo"].items(), key=lambda x: int(x[0]))
    ]
    _escribir_hoja(
        ws,
        ["Tipo", "Descripcion", "Total CSV", "Total ODS", "Coincidencias", "Cobertura (%)"],
        filas_tipo,
    )
    for fila_idx, fila in enumerate(filas_tipo, start=2):
        ws.cell(row=fila_idx, column=6).fill = PatternFill(
            "solid", fgColor=_color_por_cobertura(fila[5])
        )

    # Hoja 3: coincidencias (IDs enmascarados)
    ws = wb.create_sheet("Coincidencias")
    _escribir_hoja(
        ws,
        ["Tipo", "Descripcion", "ID (enmascarado)", "Hoja ODS"],
        [
            [c["tipo_id"], c["descripcion"], enmascarar_id(c["num_id"]), c["hoja_ods"]]
            for c in resultado["coincidencias"]
        ],
    )

    # Hoja 4: solo en CSV
    ws = wb.create_sheet("Solo en CSV")
    _escribir_hoja(
        ws,
        ["Tipo", "Descripcion", "ID (enmascarado)"],
        [
            [r["tipo_id"], r["descripcion"], enmascarar_id(r["num_id"])]
            for r in resultado["solo_en_csv"]
        ],
    )

    # Hoja 5: solo en ODS (gaps de cobertura QA)
    ws = wb.create_sheet("Solo en ODS")
    _escribir_hoja(
        ws,
        ["Tipo", "Descripcion", "ID (enmascarado)", "Hoja ODS"],
        [
            [r["tipo_id"], r["descripcion"], enmascarar_id(r["num_id"]), r.get("hoja", "")]
            for r in resultado["solo_en_ods"]
        ],
    )

    wb.save(ruta)
    return os.path.abspath(ruta)
