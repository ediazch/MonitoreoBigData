"""
informe_hojas_tecnicas.py
-------------------------
Hojas de soporte tecnico: cobertura por tipo, calidad del ODS, matriz CSV por hoja
y detalles trazables (siempre con IDs enmascarados).
"""

from typing import Any

from analisis_cruce import enmascarar_id
from informe_estilos import FMT_ENTERO, FMT_PCT, ajustar_anchos, tabla, titulo
from openpyxl.styles import PatternFill

RESULTADOS = {
    "exacta": "Coincidencia exacta",
    "id_sin_tipo": "ID en ODS sin tipo",
    "otra_hoja": "ID en otro tipo u hoja",
    "no_encontrado": "No encontrado en el ODS",
    "invalida": "Fila CSV invalida",
}
CATEGORIAS = {
    "placeholder_cero": "ID = 0",
    "tipo_sin_id": "Tipo sin ID",
    "id_sin_tipo": "ID sin tipo",
    "encabezado": "Encabezado",
    "otra": "Formato no reconocido",
}


def hoja_cobertura(wb, res: dict[str, Any]) -> None:
    """Cobertura por tipo: trazabilidad (CSV a ODS) y cobertura QA (ODS a CSV)."""
    ws = wb.create_sheet("Cobertura por tipo")
    titulo(ws, "Cobertura por tipo de documento",
           "Trazabilidad = casos CSV presentes en el ODS | Cobertura QA = IDs del ODS con caso en el CSV")
    filas = [[t["tipo"], t["descripcion"], t["hoja"], t["csv"], t["ods"], t["coincidencias"],
              t["solo_csv"], t["solo_ods"], t["trazabilidad"], t["cobertura_qa"], t["semaforo"]]
             for t in res["por_tipo"]]
    s = res["resumen"]
    filas.append(["", "TOTAL", "", s["csv_total"], s["ods_validos"], s["csv_exactas"],
                  s["csv_total"] - s["csv_exactas"], s["ods_sin_qa"], s["trazabilidad"],
                  s["cobertura_global"], ""])
    tabla(ws, 4, ["Tipo", "Descripcion", "Hoja ODS", "Casos CSV", "IDs validos ODS", "Coincidencias",
                  "Solo en CSV", "Solo en ODS", "Trazabilidad", "Cobertura QA", "Estado"],
          filas, formatos={4: FMT_ENTERO, 5: FMT_ENTERO, 6: FMT_ENTERO, 7: FMT_ENTERO,
                           8: FMT_ENTERO, 9: FMT_PCT, 10: FMT_PCT}, col_estado=11)
    ajustar_anchos(ws)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:K{4 + len(filas) - 1}"


def hoja_calidad(wb, res: dict[str, Any]) -> None:
    """Calidad por hoja del ODS."""
    ws = wb.create_sheet("Calidad ODS")
    titulo(ws, "Calidad de cada hoja del ODS",
           "Calidad = filas validas / filas candidatas (con tipo o ID; sin encabezados ni relleno)")
    filas = [[h["hoja"], h["filas_leidas"], h["candidatas"], h["validas"], h["placeholder_cero"],
              h["tipo_sin_id"], h["id_sin_tipo"], h["otra"], h["duplicados"], h["notacion_cientifica"],
              h["columnas_tipicas"], "Si" if h["estructura_distinta"] else "No",
              h["calidad_pct"], h["semaforo"]] for h in res["hojas_calidad"]]
    tabla(ws, 4, ["Hoja", "Filas leidas", "Candidatas", "Validas", "ID = 0", "Tipo sin ID", "ID sin tipo",
                  "Otras", "Duplicados", "Notacion cientifica", "Columnas tipicas", "Estructura distinta",
                  "Calidad", "Estado"], filas,
          formatos={13: FMT_PCT}, col_estado=14)
    ajustar_anchos(ws)
    ws.freeze_panes = "B5"


def hoja_matriz(wb, res: dict[str, Any]) -> None:
    """Matriz: casos del CSV por tipo (filas) segun la hoja del ODS donde aparece el ID (columnas)."""
    ws = wb.create_sheet("Matriz CSV por hoja")
    titulo(ws, "Matriz de cruce: tipo del CSV contra hoja del ODS",
           "Cada caso del CSV se busca en las 10 hojas. Lo esperado es una diagonal; "
           "valores fuera de ella indican IDs en otra hoja.")
    matriz = res["matriz"]
    columnas = list(next(iter(matriz.values())).keys()) if matriz else []
    filas = [[f"Tipo {t}"] + [matriz[t][c] for c in columnas] for t in matriz]
    fin = tabla(ws, 4, ["Tipo del CSV"] + columnas, filas,
                formatos={j: FMT_ENTERO for j in range(2, len(columnas) + 2)})
    for i, t in enumerate(matriz):
        for j, c in enumerate(columnas):
            celda = ws.cell(row=5 + i, column=j + 2)
            if celda.value and c == "No encontrado":
                celda.fill = PatternFill("solid", fgColor="FFC7CE")
            elif celda.value and c.split("_")[0] == t:
                celda.fill = PatternFill("solid", fgColor="C6EFCE")
            elif celda.value:
                celda.fill = PatternFill("solid", fgColor="FFEB9C")
    ajustar_anchos(ws, minimo=11, maximo=22)
    ws.freeze_panes = "B5"


def hoja_detalle_csv(wb, res: dict[str, Any]) -> None:
    """Un renglon por caso del CSV, con su resultado y la fila de origen en ambos archivos."""
    ws = wb.create_sheet("Detalle CSV")
    titulo(ws, "Detalle de cada caso del CSV", "Filtrar la columna Resultado para ver las excepciones.")
    filas = [[d["fila_csv"], d["tipo"], enmascarar_id(d["num_id"]), RESULTADOS.get(d["resultado"], d["resultado"]),
              d["hoja_ods"], d["fila_ods"], ", ".join(d["hojas_encontradas"]), d["observacion"]]
             for d in res["detalle_csv"]]
    tabla(ws, 4, ["Fila CSV", "Tipo", "ID (enmascarado)", "Resultado", "Hoja ODS", "Fila ODS",
                  "Hojas donde aparece", "Observacion"], filas)
    ajustar_anchos(ws, maximo=48)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{4 + len(filas)}"


def hoja_sin_qa(wb, res: dict[str, Any]) -> None:
    """IDs validos del ODS que no tienen caso en el CSV."""
    ws = wb.create_sheet("ODS sin caso QA")
    titulo(ws, "IDs del ODS sin caso de prueba en el CSV",
           "Lista de trabajo para decidir el alcance de pruebas.")
    filas = [[r["hoja"], r["fila"], r["tipo"], enmascarar_id(r["num_id"])] for r in res["ods_sin_qa"]]
    tabla(ws, 4, ["Hoja ODS", "Fila ODS", "Tipo", "ID (enmascarado)"], filas)
    ajustar_anchos(ws)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:D{4 + len(filas)}"


def hoja_excluidos(wb, res: dict[str, Any]) -> None:
    """Filas del ODS que no se usaron en el cruce y el motivo."""
    ws = wb.create_sheet("Excluidos ODS")
    titulo(ws, "Filas del ODS excluidas del analisis",
           "Se listan para que el duenio del ODS pueda revisarlas en el origen.")
    filas = [[e["hoja"], e["fila"], CATEGORIAS.get(e["categoria"], e["categoria"]),
              e["tipo"], enmascarar_id(e["num_id"]) if e["num_id"] else "", e["celdas"]]
             for e in res["ods_excluidos"]]
    tabla(ws, 4, ["Hoja ODS", "Fila ODS", "Motivo", "Tipo", "ID (enmascarado)", "Columnas"], filas)
    ajustar_anchos(ws)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{4 + max(len(filas), 1)}"
