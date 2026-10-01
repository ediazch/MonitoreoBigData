"""
informe_estilos.py
------------------
Estilos y utilidades compartidas por las hojas del informe Excel.
"""

from typing import Any, Iterable, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

AZUL = "1F3864"
AZUL_CLARO = "D9E2F3"
GRIS_CLARO = "F2F2F2"

# (color de fondo, color de texto) por estado de semaforo
SEMAFORO = {
    "VERDE": ("C6EFCE", "006100"),
    "AMARILLO": ("FFEB9C", "7F6000"),
    "ROJO": ("FFC7CE", "9C0006"),
    "N/A": ("EDEDED", "595959"),
}

FMT_PCT = '0.00"%"'
FMT_ENTERO = "#,##0"

_LADO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)


def titulo(ws, texto: str, subtitulo: Optional[str] = None) -> None:
    """Titulo grande en A1 y subtitulo opcional en A2."""
    ws["A1"] = texto
    ws["A1"].font = Font(size=16, bold=True, color=AZUL)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(size=10, italic=True, color="595959")


def seccion(ws, fila: int, texto: str, ncols: int) -> int:
    """Franja azul con el titulo de una seccion. Retorna la fila siguiente."""
    for col in range(1, ncols + 1):
        ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=AZUL)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(size=12, bold=True, color="FFFFFF")
    celda.alignment = Alignment(vertical="center")
    ws.row_dimensions[fila].height = 20
    return fila + 1


def estilo_encabezado(celda) -> None:
    """Estilo comun de encabezados de tabla."""
    celda.font = Font(bold=True, color=AZUL)
    celda.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    celda.border = BORDE
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def pintar_estado(celda, estado: Any) -> None:
    """Colorea una celda segun el estado del semaforo."""
    fondo, texto = SEMAFORO.get(str(estado), SEMAFORO["N/A"])
    celda.fill = PatternFill("solid", fgColor=fondo)
    celda.font = Font(bold=True, color=texto)
    celda.alignment = Alignment(horizontal="center", vertical="center")


def tabla(ws, fila: int, encabezados: list[str], filas: Iterable[list[Any]],
          formatos: Optional[dict[int, str]] = None, col_estado: Optional[int] = None) -> int:
    """
    Escribe encabezado y filas desde `fila`.

    Args:
        formatos: {numero_de_columna: formato_numerico} aplicado a valores numericos.
        col_estado: numero de columna cuyo valor es un estado de semaforo.

    Returns:
        Numero de la primera fila libre despues de la tabla.
    """
    for j, texto in enumerate(encabezados, start=1):
        estilo_encabezado(ws.cell(row=fila, column=j, value=texto))
    ws.row_dimensions[fila].height = 32

    siguiente = fila + 1
    for datos in filas:
        for j, valor in enumerate(datos, start=1):
            celda = ws.cell(row=siguiente, column=j, value=valor)
            celda.border = BORDE
            celda.alignment = Alignment(vertical="top",
                                        wrap_text=isinstance(valor, str) and len(valor) > 45)
            if formatos and j in formatos and isinstance(valor, (int, float)):
                celda.number_format = formatos[j]
            if col_estado == j:
                pintar_estado(celda, valor)
        siguiente += 1
    return siguiente


def ajustar_anchos(ws, minimo: int = 9, maximo: int = 60) -> None:
    """Ajusta el ancho de cada columna al contenido (con tope)."""
    anchos: dict[str, int] = {}
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is None:
                continue
            largo = max(len(parte) for parte in str(celda.value).split("\n"))
            anchos[celda.column_letter] = max(anchos.get(celda.column_letter, 0), largo)
    for letra, ancho in anchos.items():
        ws.column_dimensions[letra].width = max(minimo, min(ancho + 2, maximo))
