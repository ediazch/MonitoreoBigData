"""
informe_dashboard.py
--------------------
Hoja Dashboard: tarjetas de KPI con semaforo y graficos nativos de Excel.
Los graficos leen de una tabla de datos visible en la misma hoja.
"""

from typing import Any

from informe_estilos import (AZUL, FMT_PCT, SEMAFORO, ajustar_anchos, estilo_encabezado,
                             seccion, titulo)
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_FILA_DATOS = 52
_LADO = Side(style="medium", color="FFFFFF")


def _tarjeta(ws, fila: int, col: int, kpi: dict[str, Any]) -> None:
    """Tarjeta de 3 filas x 2 columnas: nombre, valor y estado."""
    fondo, texto = SEMAFORO.get(kpi["semaforo"], SEMAFORO["N/A"])
    relleno = PatternFill("solid", fgColor=fondo)
    for r in range(fila, fila + 3):
        for c in (col, col + 1):
            ws.cell(row=r, column=c).fill = relleno
            ws.cell(row=r, column=c).border = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)

    ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + 1)
    ws.merge_cells(start_row=fila + 1, start_column=col, end_row=fila + 1, end_column=col + 1)
    ws.merge_cells(start_row=fila + 2, start_column=col, end_row=fila + 2, end_column=col + 1)

    n = ws.cell(row=fila, column=col, value=kpi["nombre"])
    n.font = Font(size=9, bold=True, color=texto)
    n.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    valor = kpi["valor"]
    if isinstance(valor, float):
        texto_valor = f"{valor:.1f}%" if kpi["unidad"] == "%" else f"{valor:,.1f}"
    else:
        texto_valor = f"{valor:,}" if valor is not None else "n/d"
    v = ws.cell(row=fila + 1, column=col, value=texto_valor)
    v.font = Font(size=22, bold=True, color=texto)
    v.alignment = Alignment(horizontal="center", vertical="center")

    e = ws.cell(row=fila + 2, column=col, value=kpi["semaforo"] if kpi["semaforo"] != "N/A" else kpi["unidad"])
    e.font = Font(size=9, bold=True, color=texto)
    e.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[fila].height = 28
    ws.row_dimensions[fila + 1].height = 36
    ws.row_dimensions[fila + 2].height = 18


def _tabla_datos(ws, res: dict[str, Any]) -> tuple[int, int]:
    """Tabla de datos que alimenta los graficos. Retorna (fila_inicio, fila_fin)."""
    fila = seccion(ws, _FILA_DATOS - 1, "Datos del dashboard (alimentan los graficos)", 8)
    for j, t in enumerate(["Tipo", "Trazabilidad %", "Cobertura QA %", "Calidad ODS %"], start=1):
        estilo_encabezado(ws.cell(row=fila, column=j, value=t))
    calidad = {h["tipo"]: h["calidad_pct"] for h in res["hojas_calidad"]}
    inicio = fila + 1
    for i, t in enumerate(res["por_tipo"]):
        r = inicio + i
        ws.cell(row=r, column=1, value=f"{t['tipo']} {t['descripcion'][:18]}")
        for j, v in enumerate([t["trazabilidad"], t["cobertura_qa"], calidad.get(t["tipo"])], start=2):
            c = ws.cell(row=r, column=j, value=v)
            c.number_format = FMT_PCT
    return inicio, inicio + len(res["por_tipo"]) - 1


def _grafico_barras(ws, inicio: int, fin: int, columnas: list[int], titulo_g: str, ancla: str,
                    colores: list[str]) -> None:
    """Barras horizontales agrupadas por tipo de documento."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.title = titulo_g
    ch.style = 10
    ch.height, ch.width = 8.5, 16
    for col, color in zip(columnas, colores):
        ref = Reference(ws, min_col=col, min_row=inicio - 1, max_row=fin)
        ch.add_data(ref, titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
    ch.set_categories(Reference(ws, min_col=1, min_row=inicio, max_row=fin))
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 100
    ch.y_axis.title = "%"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.legend.position = "b"
    ws.add_chart(ch, ancla)


def _grafico_rosca(ws, res: dict[str, Any], ancla: str) -> None:
    """Rosca: casos del CSV con y sin referencia exacta en el ODS."""
    s = res["resumen"]
    fila = _FILA_DATOS + len(res["por_tipo"]) + 3
    estilo_encabezado(ws.cell(row=fila, column=1, value="Casos CSV"))
    estilo_encabezado(ws.cell(row=fila, column=2, value="Cantidad"))
    ws.cell(row=fila + 1, column=1, value="Con referencia exacta")
    ws.cell(row=fila + 1, column=2, value=s["csv_exactas"])
    ws.cell(row=fila + 2, column=1, value="Sin referencia")
    ws.cell(row=fila + 2, column=2, value=s["csv_total"] - s["csv_exactas"])

    ch = DoughnutChart()
    ch.title = "Casos del CSV con referencia en el ODS"
    ch.height, ch.width = 8.5, 11
    ch.add_data(Reference(ws, min_col=2, min_row=fila, max_row=fila + 2), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=fila + 1, max_row=fila + 2))
    for idx, color in enumerate(["70AD47", "C00000"]):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = color
        ch.series[0].dPt.append(pt)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showPercent = False
    ch.legend.position = "b"
    ws.add_chart(ch, ancla)


def hoja_dashboard(wb, res: dict[str, Any]) -> None:
    """Dashboard: tarjetas de KPI y tres graficos. Se coloca despues del resumen ejecutivo."""
    ws = wb.create_sheet("Dashboard", 1)
    titulo(ws, "Dashboard - Cruce CSV vs ODS", f"Generado: {res['generado']}")
    ws.sheet_view.showGridLines = False

    seccion(ws, 3, "Indicadores clave", 8)
    for i, kpi in enumerate(res["kpis"][:8]):
        _tarjeta(ws, 4 if i < 4 else 8, 1 + (i % 4) * 2, kpi)

    seccion(ws, 12, "Cobertura y trazabilidad por tipo de documento", 8)
    inicio, fin = _tabla_datos(ws, res)
    _grafico_barras(ws, inicio, fin, [2, 3], "Trazabilidad y cobertura QA por tipo (%)", "A13",
                    ["2E75B6", "ED7D31"])
    _grafico_rosca(ws, res, "F13")
    _grafico_barras(ws, inicio, fin, [4], "Calidad de cada hoja del ODS (%)", "A31", ["70AD47"])

    for letra in "ABCDEFGH":
        ws.column_dimensions[letra].width = 17
    ws.column_dimensions["A"].width = 30
