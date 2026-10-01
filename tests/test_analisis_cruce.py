"""
test_analisis_cruce.py
----------------------
Pruebas del cruce CSV vs ODS con datos sinteticos (nunca datos reales).
Cubren: lectura de ODS (notacion cientifica, filas repetidas, XXE), clasificacion
de filas, cruce contra todas las hojas, KPIs y semaforo, y el Excel generado.
"""

import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from analisis_cruce import (analizar_datos, clasificar_fila_ods, enmascarar_id, pct,
                            semaforo)
from limpieza import leer_ods_detallado

UMBRALES = {
    "cobertura_qa": {"verde": 90.0, "amarillo": 70.0},
    "trazabilidad": {"verde": 99.0, "amarillo": 95.0},
    "calidad_ods": {"verde": 95.0, "amarillo": 85.0},
}


def fila_ods(hoja, fila, *celdas, notas=()):
    return {"hoja": hoja, "fila": fila, "celdas": list(celdas), "notas": list(notas)}


def _ods_sintetico():
    return [
        fila_ods("1_CC", 1, "1", "1000000001", "10"),
        fila_ods("1_CC", 2, "1", "1000000002", "10"),
        fila_ods("1_CC", 3, "1", "1000000003", "10"),
        fila_ods("1_CC", 4, "1", "0", "1"),
        fila_ods("4_TI", 1, "4", "4000000001", "10"),
        fila_ods("4_TI", 2, "4", "", "10"),
        fila_ods("13_PPT", 1, "tipoID", "NumID"),
        fila_ods("13_PPT", 2, "", "1300000001"),
    ]


def _csv_sintetico():
    return [(1, ["1", "1000000001"]), (2, ["1", "1000000002"]),
            (3, ["4", "4000000001"]), (4, ["1", "9999999999"]), (5, ["13", "1300000001"])]


class TestUtilidades:
    def test_enmascarar_id_largo(self):
        assert enmascarar_id("1069729185") == "106****185"

    def test_enmascarar_id_corto(self):
        assert enmascarar_id("12345") == "***"

    def test_pct_sin_denominador(self):
        assert pct(1, 0) is None

    def test_pct_redondea(self):
        assert pct(1, 3) == 33.33

    @pytest.mark.parametrize("valor,esperado", [
        (95.0, "VERDE"), (90.0, "VERDE"), (89.99, "AMARILLO"), (70.0, "AMARILLO"),
        (69.99, "ROJO"), (0.0, "ROJO"), (None, "N/A"),
    ])
    def test_semaforo_limites(self, valor, esperado):
        assert semaforo(valor, UMBRALES["cobertura_qa"]) == esperado


class TestClasificarFilaOds:
    @pytest.mark.parametrize("celdas,categoria", [
        (["1", "123456"], "valida"),
        (["1", "0"], "placeholder_cero"),
        (["1", ""], "tipo_sin_id"),
        (["", "123456"], "id_sin_tipo"),
        (["", ""], "relleno"),
        (["tipoID", "NumID"], "encabezado"),
        (["1", "ABC"], "otra"),
        ([], "relleno"),
    ])
    def test_categorias(self, celdas, categoria):
        assert clasificar_fila_ods(celdas)[0] == categoria

    def test_limpia_comillas_en_id(self):
        assert clasificar_fila_ods(["1", "'123456'"]) == ("valida", "1", "123456")


class TestCruce:
    def setup_method(self):
        self.res = analizar_datos(_csv_sintetico(), _ods_sintetico(), UMBRALES)
        self.por_resultado = {d["fila_csv"]: d["resultado"] for d in self.res["detalle_csv"]}

    def test_coincidencias_exactas(self):
        assert [self.por_resultado[i] for i in (1, 2, 3)] == ["exacta"] * 3

    def test_id_inexistente_no_se_encuentra(self):
        assert self.por_resultado[4] == "no_encontrado"

    def test_id_en_fila_sin_tipo_se_reporta_aparte(self):
        # La fila ODS de 13_PPT no declara tipo: no cuenta como exacta pero tampoco se pierde.
        assert self.por_resultado[5] == "id_sin_tipo"

    def test_resumen_conteos(self):
        s = self.res["resumen"]
        assert s["csv_total"] == 5
        assert s["csv_exactas"] == 3
        assert s["ods_validos"] == 4
        assert s["ods_sin_qa"] == 1

    def test_ods_sin_qa_es_el_id_no_probado(self):
        assert [r["num_id"] for r in self.res["ods_sin_qa"]] == ["1000000003"]

    def test_filas_excluidas_se_registran_con_su_fila(self):
        motivos = {(e["hoja"], e["fila"]): e["categoria"] for e in self.res["ods_excluidos"]}
        assert motivos[("1_CC", 4)] == "placeholder_cero"
        assert motivos[("4_TI", 2)] == "tipo_sin_id"
        assert motivos[("13_PPT", 2)] == "id_sin_tipo"

    def test_cruza_contra_todas_las_hojas(self):
        ods = _ods_sintetico() + [fila_ods("4_TI", 3, "4", "1000000099", "10")]
        res = analizar_datos([(1, ["1", "1000000099"])], ods, UMBRALES)
        d = res["detalle_csv"][0]
        assert d["resultado"] == "otra_hoja"
        assert d["hojas_encontradas"] == ["4_TI"]

    def test_csv_invalido_no_cuenta_en_el_total(self):
        res = analizar_datos([(1, ["x", "abc"]), (2, ["1", "1000000001"])], _ods_sintetico(), UMBRALES)
        assert res["resumen"]["csv_total"] == 1
        assert res["detalle_csv"][0]["resultado"] == "invalida"

    def test_cobertura_y_trazabilidad_son_metricas_distintas(self):
        s = self.res["resumen"]
        assert s["trazabilidad"] == pct(3, 5)
        assert s["cobertura_global"] == pct(3, 4)
        assert s["trazabilidad"] != s["cobertura_global"]

    def test_kpis_incluyen_semaforo(self):
        estados = {k["nombre"]: k["semaforo"] for k in self.res["kpis"]}
        assert estados["Trazabilidad CSV a ODS"] == "ROJO"

    def test_vacios_no_rompen(self):
        res = analizar_datos([], [], UMBRALES)
        assert res["resumen"]["trazabilidad"] is None
        assert res["resumen"]["cobertura_global"] is None


def _escribir_ods(ruta, contenido_xml):
    with zipfile.ZipFile(ruta, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", contenido_xml)


_CABECERA = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet>'
)
_PIE = "</office:spreadsheet></office:body></office:document-content>"


def _celda(texto, extra=""):
    return f'<table:table-cell {extra}><text:p>{texto}</text:p></table:table-cell>'


class TestLecturaOds:
    def test_recupera_notacion_cientifica(self, tmp_path):
        xml = (_CABECERA + '<table:table table:name="6_X"><table:table-row>'
               + _celda("6") + _celda("1,23E+09", 'office:value-type="float" office:value="1234567890"')
               + "</table:table-row></table:table>" + _PIE)
        ruta = tmp_path / "t.ods"
        _escribir_ods(ruta, xml)
        filas = leer_ods_detallado(str(ruta))
        assert filas[0]["celdas"][1] == "1234567890"
        assert filas[0]["notas"] == ["notacion_cientifica_recuperada"]

    def test_expande_filas_repetidas_con_numero_correcto(self, tmp_path):
        xml = (_CABECERA + '<table:table table:name="1_X">'
               '<table:table-row table:number-rows-repeated="3">' + _celda("1") + _celda("111111") + "</table:table-row>"
               "<table:table-row>" + _celda("1") + _celda("222222") + "</table:table-row></table:table>" + _PIE)
        ruta = tmp_path / "t.ods"
        _escribir_ods(ruta, xml)
        assert [f["fila"] for f in leer_ods_detallado(str(ruta))] == [1, 2, 3, 4]

    def test_rechaza_dtd_entity(self, tmp_path):
        xml = ('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]>' + _CABECERA[_CABECERA.index("<office"):] + _PIE)
        ruta = tmp_path / "t.ods"
        _escribir_ods(ruta, xml)
        with pytest.raises(ValueError, match="DTD"):
            leer_ods_detallado(str(ruta))


class TestInformeExcel:
    def test_genera_hojas_y_no_expone_ids(self, tmp_path):
        from informe_excel import generar_informe
        from openpyxl import load_workbook

        res = analizar_datos(_csv_sintetico(), _ods_sintetico(), UMBRALES)
        ruta = generar_informe(res, str(tmp_path))
        wb = load_workbook(ruta)
        assert wb.sheetnames == [
            "Resumen ejecutivo", "Dashboard", "Cobertura por tipo", "Calidad ODS", "Matriz CSV por hoja",
            "Detalle CSV", "ODS sin caso QA", "Excluidos ODS", "Metodologia",
        ]
        assert len(wb["Dashboard"]._charts) == 3
        ids_reales = {"1000000001", "1000000002", "1000000003", "4000000001", "1300000001", "9999999999"}
        for ws in wb.worksheets:
            for fila in ws.iter_rows():
                for c in fila:
                    assert str(c.value) not in ids_reales, f"ID completo en {ws.title}!{c.coordinate}"
